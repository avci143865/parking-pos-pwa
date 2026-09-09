"""تصدير طوابير FIFO إلى Excel (يدعم العربية بدون خطوط إضافية)."""

from __future__ import annotations

import re
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.fifo_queue import DEFAULT_TRUCK_TYPE

_INVALID_SHEET_TITLE = re.compile(r"[\\/*?:\[\]]")
_META_SHEET_TITLE = "معلومات"


def _sanitize_sheet_title(name: str, used: set[str]) -> str:
    """أسماء أوراق Excel: بدون \\ / ? * [ ] وبحد 31 حرفاً وبدون تكرار."""
    base = _INVALID_SHEET_TITLE.sub("-", (name or "طابور").strip()) or "طابور"
    base = base[:31]
    title = base
    n = 2
    while title in used:
        suffix = f" ({n})"
        title = (base[: max(1, 31 - len(suffix))] + suffix)[:31]
        n += 1
    used.add(title)
    return title


def _status_label(item: dict, allowed: int | None) -> str:
    label = item.get("status_label")
    if label:
        return str(label)
    if item.get("queue_status") == "ready_exit":
        return "قيد الخروج"
    return "في الانتظار"


def _write_queue_sheet(ws, q: dict) -> None:
    ws.sheet_view.rightToLeft = True
    truck_type = q.get("truck_type") or DEFAULT_TRUCK_TYPE
    waiting = q.get("waiting_count", 0)
    total_waiting = q.get("total_waiting")
    export_limit = q.get("export_limit")
    allowed = q.get("current_allowed_queue_number")

    if export_limit and total_waiting is not None and int(export_limit) < int(total_waiting):
        scope = f"أول {waiting} من {total_waiting}"
    else:
        scope = f"منتظر: {waiting}"
    allowed_txt = str(allowed) if allowed is not None else "—"

    ws.append([f"نوع الشاحنة: {truck_type} | {scope} | الدور المسموح: {allowed_txt}"])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=7)
    ws["A1"].font = Font(bold=True, size=12)
    ws["A1"].alignment = Alignment(horizontal="right", vertical="center")
    ws.append([])

    headers = (
        "رقم الدور",
        "اللوحة",
        "السائق",
        "الشركة",
        "النوع",
        "وقت الدخول",
        "الحالة",
    )
    ws.append(list(headers))
    header_row = ws.max_row
    fill = PatternFill("solid", fgColor="E6ECF5")
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=header_row, column=col)
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for item in q.get("items") or []:
        ws.append(
            [
                item.get("fifo_queue_number") or "—",
                item.get("license_plate") or "—",
                item.get("driver_name") or "—",
                item.get("partnership_company") or "—",
                item.get("fifo_truck_type") or truck_type,
                item.get("entered_at_display") or "—",
                _status_label(item, allowed),
            ]
        )

    widths = (8, 14, 22, 22, 14, 22, 18)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build_fifo_queues_xlsx(
    *,
    title: str,
    exported_at: datetime,
    queues: list[dict],
) -> bytes:
    wb = Workbook()
    used_titles: set[str] = {_META_SHEET_TITLE}
    if queues:
        first = queues[0]
        ws = wb.active
        ws.title = _sanitize_sheet_title(first.get("truck_type") or "طابور", used_titles)
        _write_queue_sheet(ws, first)
        for q in queues[1:]:
            sheet = wb.create_sheet(
                title=_sanitize_sheet_title(q.get("truck_type") or "طابور", used_titles)
            )
            _write_queue_sheet(sheet, q)
    else:
        ws = wb.active
        ws.title = _sanitize_sheet_title("أدوار الخروج", used_titles)
        ws.sheet_view.rightToLeft = True
        ws.append([title])
        ws.append([f"تاريخ التصدير: {exported_at.strftime('%Y-%m-%d %H:%M')} (دمشق)"])
        ws.append(["لا توجد بيانات."])

    meta = wb.create_sheet(title=_META_SHEET_TITLE, index=0)
    meta.sheet_view.rightToLeft = True
    meta.append([title])
    meta.append([f"تاريخ التصدير: {exported_at.strftime('%Y-%m-%d %H:%M')} (دمشق)"])
    meta.column_dimensions["A"].width = 48

    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio.getvalue()
